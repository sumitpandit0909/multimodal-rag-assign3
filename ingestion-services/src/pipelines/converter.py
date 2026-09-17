import subprocess
import os
import shutil
import tempfile
from pathlib import Path
import logging
from typing import Optional, List, Dict

logger = logging.getLogger(__name__)

def convert_to_pdf(input_file_path: Path, output_dir: Path) -> Path:
    """
    Converts DOCX/PPT/PPTX to PDF using headless LibreOffice with multiple filter attempts.
    """
    if input_file_path.suffix.lower() == ".pdf":
        return input_file_path

    expected_pdf = output_dir / f"{input_file_path.stem}.pdf"
    if expected_pdf.exists():
        try:
            expected_pdf.unlink()
        except Exception:
            pass

    # Strategies to try with LibreOffice
    suffix = input_file_path.suffix.lower()
    specific_filter = "pdf:impress_pdf_Export" if suffix in [".ppt", ".pptx"] else "pdf:writer_pdf_Export"

    attempts = [
        # Attempt 1: Standard conversion
        [
            "libreoffice", "--headless", "--invisible", "--nodefault", 
            "--nofirststartwizard", "--nolockcheck", "--nologo", "--norestore",
            "--convert-to", "pdf",
            "--outdir", str(output_dir),
            str(input_file_path)
        ],
        # Attempt 2: Explicit application export filter
        [
            "libreoffice", "--headless", "--invisible", "--nodefault",
            "--nofirststartwizard", "--nolockcheck", "--nologo", "--norestore",
            "--convert-to", specific_filter,
            "--outdir", str(output_dir),
            str(input_file_path)
        ]
    ]

    last_error = ""
    for cmd in attempts:
        logger.info(f"Running LibreOffice conversion command: {' '.join(cmd)}")
        try:
            result = subprocess.run(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=120
            )
            if expected_pdf.exists() and expected_pdf.stat().st_size > 0:
                logger.info(f"Successfully converted {input_file_path.name} to {expected_pdf.name}")
                return expected_pdf
            last_error = result.stderr or result.stdout
        except Exception as e:
            last_error = str(e)
            logger.warning(f"LibreOffice command attempt failed: {e}")

    # If expected_pdf still not found
    if not expected_pdf.exists() or expected_pdf.stat().st_size == 0:
        raise FileNotFoundError(
            f"Expected PDF output not found at {expected_pdf}. LibreOffice output:\n{last_error}"
        )

    return expected_pdf

def parse_pptx_fallback(file_path: Path) -> List[Dict]:
    """
    Fallback parser for PPTX presentations using python-pptx when headless rendering fails.
    Extracts structured content per slide.
    """
    try:
        from pptx import Presentation
    except ImportError:
        logger.error("python-pptx is not installed.")
        return []

    prs = Presentation(file_path)
    slides_data = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_text_blocks = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_text_blocks.append(line)
            elif shape.has_table:
                table_rows = []
                for row in shape.table.rows:
                    row_vals = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(row_vals))
                if table_rows:
                    slide_text_blocks.append("\n" + "\n".join(table_rows) + "\n")

        content = "\n".join(slide_text_blocks).strip()
        if not content:
            content = f"[Slide {idx} containing diagrams/graphics]"

        slides_data.append({
            "page_number": idx,
            "text": content,
            "title": slide_text_blocks[0] if slide_text_blocks else f"Slide {idx}"
        })

    return slides_data
